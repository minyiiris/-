import streamlit as st
import pdfplumber
import re
import openpyxl
import io

st.set_page_config(page_title="報單自動化工具", layout="centered")

st.title("📄 進口報單 PDF 自動轉 Excel 工具")
st.write("請依序上傳您的「範本.xlsx」以及所有需要處理的「PDF 報單檔案」，系統會自動處理並產出結果檔案。")

# 1. 上傳區塊
template_file = st.file_uploader("第一步：請上傳 Excel 範本 (例如：範本.xlsx)", type=["xlsx"])
pdf_files = st.file_uploader("第二步：請上傳 PDF 報單檔案 (可一次全選拖曳上傳)", type=["pdf"], accept_multiple_files=True)

if st.button("🚀 開始處理並產生 Excel", type="primary"):
    if not template_file:
        st.error("請先上傳 Excel 範本！")
    elif not pdf_files:
        st.error("請至少上傳一個 PDF 檔案！")
    else:
        with st.spinner("正在自動識別資料中，請稍候..."):
            # 依據檔名自動排序
            pdf_files = sorted(pdf_files, key=lambda x: x.name)
            all_data = []
            
            for pdf_file in pdf_files:
                data = {'報關日期': '', '提單號碼': '', '報關編號': '', 'CCC_Code': '', '數量': '', '淨重': '', '貨物名稱': ''}
                with pdfplumber.open(pdf_file) as pdf:
                    text = "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
                    
                    # 報關編號
                    m_decl = re.search(r'CX\s*[0-9A-Z]{10,12}', text)
                    if m_decl: data['報關編號'] = m_decl.group(0)
                    
                    # 分提單號碼
                    m_awb = re.search(r'分提單號碼(?:\(9\))?\s*(\d+)', text)
                    if m_awb: data['提單號碼'] = m_awb.group(1)
                    
                    # 報關日期 (轉西元年)
                    m_date = re.search(r'報關日期(?:\(15\))?\s*(\d{7})', text)
                    if m_date: 
                        r_date = m_date.group(1)
                        data['報關日期'] = f"{int(r_date[:3]) + 1911}/{r_date[3:5]}/{r_date[5:]}"
                        
                    # CCC Code (自動換行)
                    ccc_matches = re.findall(r'(\d{4}\.\d{2}\.\d{2}\.\d{2}\-\d)', text)
                    if ccc_matches:
                        data['CCC_Code'] = "\n".join([c.replace(".", "").replace("-", "") for c in ccc_matches])
                        
                    # 數量與淨重
                    m_total = re.search(r'TOTAL:\s*([\d\.]+)\s*[A-Z]+\s*(\d+)\s*[A-Z]+', text)
                    if m_total:
                        data['淨重'] = float(m_total.group(1))
                        data['數量'] = int(m_total.group(2))
                        
                    # 貨物名稱
                    items_desc = []
                    parts = text.split("VIET NAM") 
                    for p in parts[1:]:
                        m_desc = re.search(r'(.*?)\n\d{4}\.\d{2}\.\d{2}\.\d{2}\-\d', p, re.DOTALL)
                        if m_desc:
                            items_desc.append(m_desc.group(1).strip().replace('\n', ' '))
                    if items_desc:
                        data['貨物名稱'] = "\n".join(items_desc)

                    all_data.append(data)
                    
            # 寫入 Excel 範本
            wb = openpyxl.load_workbook(template_file)
            ws = wb.active
            
            # 尋找標題列位置並建立欄位對應 (假設標題在第4列)
            col_map = {}
            for col in range(1, ws.max_column + 1):
                val = ws.cell(row=4, column=col).value 
                if val: col_map[val.strip()] = col
                
            start_row = 5
            for idx, row_data in enumerate(all_data):
                current_row = start_row + idx
                for key, val in row_data.items():
                    if key in col_map:
                        cell = ws.cell(row=current_row, column=col_map[key], value=val)
                        if key in ['CCC_Code', '貨物名稱']: 
                            cell.alignment = openpyxl.styles.Alignment(wrap_text=True)
                            
            # 將 Excel 存入記憶體準備下載
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            st.success(f"🎉 處理完成！共處理 {len(pdf_files)} 筆資料，請點擊下方按鈕下載檔案。")
            
            st.download_button(
                label="📥 下載自動填入結果.xlsx",
                data=output,
                file_name="自動填入結果.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
